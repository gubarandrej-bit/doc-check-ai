"""Парсинг загруженных файлов в структурированные данные.

Важно: при отсутствии данных НИЧЕГО не придумывается.
Если формат не поддерживается или данные не извлекаются — возвращается None
и проверка помечается как не проведённая с указанием причины.
"""
import os
from typing import Any


def detect_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if ext in ("xls", "xlsx"):
        return "xls"
    if ext in ("doc", "docx"):
        return "doc"
    if ext == "pdf":
        return "pdf"
    if ext == "dwg":
        return "dwg"
    return ext or "unknown"


def parse_xls(path: str) -> dict | None:
    """Возвращает {'sheets': {name: [rows]} } или None."""
    try:
        import openpyxl
    except Exception:
        return None
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        out = {}
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    rows.append([("" if c is None else str(c)) for c in row])
            out[ws.title] = rows
        return {"sheets": out} if out else None
    except Exception:
        return None


def parse_doc(path: str) -> dict | None:
    """Возвращает {'text': ...} или None."""
    try:
        import docx
    except Exception:
        return None
    try:
        d = docx.Document(path)
        paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
        tables = []
        for t in d.tables:
            rows = []
            for r in t.rows:
                rows.append([c.text for c in r.cells])
            tables.append(rows)
        return {"text": "\n".join(paragraphs), "tables": tables} if (paragraphs or tables) else None
    except Exception:
        return None


# Ключевые слова страницы, по которым таблица относится к спецификации,
# кабельному журналу и т.д. Движок различает листы по подстроке в имени.
_PDF_TABLE_HINTS = ("специф", "материал", "кабель", "журнал", "нагруз",
                    "источник", "аккумулятор", "батаре", "оборуд")


def _pdf_tables(path: str) -> dict[str, list]:
    """Таблицы из PDF в том же виде, что и листы xls: {имя: строки}.

    Имя листа берётся из текста страницы — по нему движок решает, это
    спецификация, кабельный журнал или что-то ещё. pdfplumber находит
    таблицы по их рамкам; если рамок нет, таблиц не будет и это честно
    отражается в проверке как «структура не распознана».
    """
    try:
        import pdfplumber
    except Exception:
        return {}
    out: dict[str, list] = {}
    try:
        with pdfplumber.open(path) as pdf:
            for pno, page in enumerate(pdf.pages, 1):
                try:
                    page_text = (page.extract_text() or "").lower()
                    tables = page.extract_tables() or []
                except Exception:
                    continue
                for i, tbl in enumerate(tables, 1):
                    rows = [[(c or "").strip() for c in row] for row in tbl]
                    rows = [r for r in rows if any(c for c in r)]
                    if not rows:
                        continue
                    hint = next((h for h in _PDF_TABLE_HINTS if h in page_text), "")
                    out[f"стр.{pno} {hint} {i}".strip()] = rows
    except Exception:
        return out
    return out


def parse_pdf(path: str) -> dict | None:
    """Текст и таблицы из PDF. DWG-чертежи текстом не извлекаются."""
    try:
        from pypdf import PdfReader
    except Exception:
        return None
    try:
        reader = PdfReader(path)
        text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
        sheets = _pdf_tables(path)
        if not text.strip() and not sheets:
            return None
        out = {"text": text, "pages": len(reader.pages)}
        # Таблицы отдаём как sheets, чтобы PDF-спецификация питала те же
        # проверки, что и xls: иначе проект с исходниками в PDF оставался бы
        # без единой выполненной детерминированной проверки.
        if sheets:
            out["sheets"] = sheets
        return out
    except Exception:
        return None


def parse_file(path: str, ftype: str) -> dict | None:
    if ftype == "xls":
        return parse_xls(path)
    if ftype == "doc":
        return parse_doc(path)
    if ftype == "pdf":
        return parse_pdf(path)
    if ftype == "dwg":
        # DWG-файлы не могут быть расшифрованы без специализированного
        # ПО (например, ODA File Converter). Данные не извлекаются.
        return None
    return None